#!/usr/bin/env python3
"""Convierte archivos de Office (docx, pptx, xlsx) a Markdown con imagenes reales.

Motores (nativos, sin markitdown):
  docx  -> mammoth: texto + tablas + imagenes extraidas a disco.
  pptx  -> python-pptx: texto por slide, tablas, notas del orador e imagenes.
  xlsx  -> openpyxl: una tabla Markdown por hoja.
  otros -> markitdown como respaldo (html, csv, json, xml, epub, msg, zip).

Uso:
  py office_a_markdown.py examen.docx
  py office_a_markdown.py clase.pptx -o clase.md --imagenes Figuras/clase
  py office_a_markdown.py notas.xlsx

Salida:
  <salida>.md y, si hay imagenes, la carpeta indicada (por defecto <salida>_imagenes).
"""
import argparse
import re
import sys
from datetime import datetime
from pathlib import Path

CONTROL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def limpiar(texto):
    return CONTROL.sub("", texto or "").replace("\r\n", "\n").strip()


def tabla_markdown(filas):
    filas = [[limpiar(celda).replace("|", "\\|").replace("\n", " ") for celda in fila] for fila in filas]
    filas = [fila for fila in filas if any(celda for celda in fila)]
    if not filas:
        return ""
    ancho = max(len(fila) for fila in filas)
    filas = [fila + [""] * (ancho - len(fila)) for fila in filas]
    lineas = ["| " + " | ".join(filas[0]) + " |", "| " + " | ".join(["---"] * ancho) + " |"]
    lineas += ["| " + " | ".join(fila) + " |" for fila in filas[1:]]
    return "\n".join(lineas)


def docx_a_markdown(ruta, carpeta_imagenes, prefijo):
    import mammoth
    from markdownify import markdownify

    estado = {"n": 0}

    def guardar_imagen(image):
        with image.open() as contenido:
            datos = contenido.read()
        extension = (image.content_type or "image/png").split("/")[-1].replace("jpeg", "jpg")
        estado["n"] += 1
        nombre = f"{prefijo}-imagen-{estado['n']:02d}.{extension}"
        carpeta_imagenes.mkdir(parents=True, exist_ok=True)
        (carpeta_imagenes / nombre).write_bytes(datos)
        return {"src": f"{carpeta_imagenes.name}/{nombre}"}

    with open(ruta, "rb") as archivo:
        resultado = mammoth.convert_to_html(
            archivo,
            convert_image=mammoth.images.img_element(guardar_imagen),
        )
    return markdownify(resultado.value, heading_style="ATX"), estado["n"]


def pptx_a_markdown(ruta, carpeta_imagenes, prefijo):
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    presentacion = Presentation(str(ruta))
    bloques = []
    total_imagenes = 0

    for indice, slide in enumerate(presentacion.slides, start=1):
        titulo = ""
        if slide.shapes.title is not None:
            titulo = limpiar(slide.shapes.title.text)
        encabezado = f"## Slide {indice}" + (f": {titulo}" if titulo else "")
        partes = [encabezado]

        for forma in slide.shapes:
            if forma.shape_type == MSO_SHAPE_TYPE.PICTURE:
                total_imagenes += 1
                extension = forma.image.ext
                nombre = f"{prefijo}-slide{indice:02d}-img{total_imagenes:02d}.{extension}"
                carpeta_imagenes.mkdir(parents=True, exist_ok=True)
                (carpeta_imagenes / nombre).write_bytes(forma.image.blob)
                alt = limpiar(forma.name) or nombre
                partes.append(f"![{alt}]({carpeta_imagenes.name}/{nombre})")
                continue

            if forma.has_table:
                filas = [[celda.text for celda in fila.cells] for fila in forma.table.rows]
                tabla = tabla_markdown(filas)
                if tabla:
                    partes.append(tabla)
                continue

            if not forma.has_text_frame:
                continue
            if slide.shapes.title is not None and forma == slide.shapes.title:
                continue
            for parrafo in forma.text_frame.paragraphs:
                texto = limpiar(parrafo.text)
                if not texto:
                    continue
                nivel = min(parrafo.level or 0, 4)
                partes.append("  " * nivel + "- " + texto if nivel else texto)

        if slide.has_notes_slide:
            notas = limpiar(slide.notes_slide.notes_text_frame.text)
            if notas:
                partes.append("### Notas del orador\n\n" + notas)

        bloques.append("\n\n".join(partes))

    return "\n\n".join(bloques), total_imagenes


def xlsx_a_markdown(ruta, carpeta_imagenes, prefijo):
    import openpyxl

    libro = openpyxl.load_workbook(str(ruta), data_only=True)
    bloques = []
    for hoja in libro.worksheets:
        filas = []
        for fila in hoja.iter_rows(values_only=True):
            filas.append(["" if valor is None else str(valor) for valor in fila])
        tabla = tabla_markdown(filas)
        bloques.append(f"## {hoja.title}\n\n" + (tabla if tabla else "(hoja vacia)"))
    return "\n\n".join(bloques), 0


def respaldo_markitdown(ruta):
    from markitdown import MarkItDown

    return MarkItDown().convert(str(ruta)).text_content, 0


def main():
    parser = argparse.ArgumentParser(description="Convierte docx, pptx, xlsx y otros a Markdown.")
    parser.add_argument("archivo", type=Path, help="Archivo de entrada")
    parser.add_argument("-o", "--salida", type=Path, default=None, help="Archivo .md de salida")
    parser.add_argument("--imagenes", type=Path, default=None, help="Carpeta para las imagenes extraidas")
    parser.add_argument("--titulo", default=None, help="Titulo de la nota")
    args = parser.parse_args()

    for flujo in (sys.stdout, sys.stderr):
        try:
            flujo.reconfigure(encoding="utf-8", errors="replace")
        except (AttributeError, ValueError):
            pass

    origen = args.archivo
    if not origen.is_file():
        print(f"Error: no existe el archivo '{origen}'.", file=sys.stderr)
        return 2

    salida = args.salida or origen.with_suffix(".md")
    carpeta_imagenes = args.imagenes or salida.with_name(salida.stem + "_imagenes")
    prefijo = origen.stem.replace(" ", "-")
    extension = origen.suffix.lower()

    if extension == ".docx":
        markdown, total = docx_a_markdown(origen, carpeta_imagenes, prefijo)
    elif extension == ".pptx":
        markdown, total = pptx_a_markdown(origen, carpeta_imagenes, prefijo)
    elif extension in {".xlsx", ".xlsm"}:
        markdown, total = xlsx_a_markdown(origen, carpeta_imagenes, prefijo)
    else:
        try:
            markdown, total = respaldo_markitdown(origen)
        except Exception as exc:
            print(f"Error: no se pudo convertir '{origen.name}': {type(exc).__name__}: {exc}", file=sys.stderr)
            return 1

    if not markdown or not markdown.strip():
        print(f"Error: '{origen.name}' no devolvio contenido.", file=sys.stderr)
        return 1

    titulo = args.titulo or origen.stem
    cabecera = (
        f"# {titulo}\n\n"
        f"Fuente: `{origen.name}` | Generado: {datetime.now().strftime('%Y-%m-%d %H:%M')} | "
        f"Imagenes: {total}\n\n"
    )
    salida.parent.mkdir(parents=True, exist_ok=True)
    salida.write_text(cabecera + markdown.strip() + "\n", encoding="utf-8")
    print(f"Guardado: {salida}")
    if total:
        print(f"Imagenes: {total} en {carpeta_imagenes}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
